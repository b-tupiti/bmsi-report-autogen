from openpyxl.chart import (
    LineChart,
    ScatterChart,
    Reference,
    Series
)
from openpyxl.chart.axis import DateAxis

def graph_fca(ws,_len):
    c1 = LineChart()
    c1.title = "First Call Activation"
    c1.y_axis.title = "Number of FCA"
    c1.x_axis.title = "Date"

    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None

    data = Reference(ws, min_col=2, min_row=3, max_row=_len)
    c1.add_data(data)
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    c1.width = 20

    # Style the lines
    s1 = c1.series[0]
    s1.marker.symbol = "circle"
    s1.marker.graphicalProperties.solidFill = "FF0000" # Marker filling
    s1.marker.graphicalProperties.line.solidFill = "DDEBF7" # Marker outline
    s1.graphicalProperties.line.noFill = False

    ws.add_chart(c1, "D3")
    return ws

def graph_ass(ws,_len):
    c1 = LineChart()
    c1.title = "Daily Active Customer Count"
    c1.x_axis.title = "Date"

    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None

    data = Reference(ws, min_col=2, min_row=3, max_row=_len)
    c1.add_data(data)
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    c1.width = 20

    # Style the lines
    s1 = c1.series[0]
    s1.marker.symbol = "circle"
    s1.marker.graphicalProperties.solidFill = "FF0000" # Marker filling
    s1.marker.graphicalProperties.line.solidFill = "DDEBF7" # Marker outline
    s1.graphicalProperties.line.noFill = False

    ws.add_chart(c1, "D3")
    
    return ws

def graph_srev(ws, _len):
    
    # create chart
    c1 = LineChart()
    c1.title = 'SI Summary Revenue'
    c1.y_axis.title = 'Total Revenue (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=10, min_row=3, max_row=_len)
    c1.add_data(data)
    
    
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    #-----------
    
    # create chart
    c2 = LineChart()
    c2.title = 'Revenue Item Comparision'
    c2.y_axis.title = 'Total Revenue (SBD)'
    c2.x_axis.title = 'Date'
    
    # customize chart
    c2.style=2
    c2.y_axis.crossAx = 500
    c2.x_axis = DateAxis(crossAx=100)
    c2.x_axis.number_format = 'yyyy-mm-dd'
    c2.x_axis.majorTimeUnit = "days"
    c2.legend = None
    
    
    # y axis data
    data = Reference(ws,min_col=9,min_row=3,max_row=_len)
    c2.add_data(data)
    
    # y axis data
    data = Reference(ws,min_col=8,min_row=3,max_row=_len)
    c2.add_data(data)
    
    # y axis data
    data = Reference(ws,min_col=7,min_row=3,max_row=_len)
    c2.add_data(data)
    
    # y axis data
    data = Reference(ws,min_col=6,min_row=3,max_row=_len)
    c2.add_data(data)
    
    # y axis data
    data = Reference(ws,min_col=5,min_row=3,max_row=_len)
    c2.add_data(data)
    
    # y axis data
    data = Reference(ws,min_col=4,min_row=3,max_row=_len)
    c2.add_data(data)
    
    # y axis data
    data = Reference(ws,min_col=3,min_row=3,max_row=_len)
    c2.add_data(data)
    
    # y axis data
    data = Reference(ws,min_col=2,min_row=3,max_row=_len)
    c2.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c2.set_categories(dates)
    
    # Style chart
    s2 = c2.series[0]
    s2.marker.symbol = 'circle'
    s2.marker.graphicalProperties.solidFill = 'FF0000' 
    s2.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s2.graphicalProperties.line.noFill = False
    
    #-----------
    
    # add chart to worksheet
    ws.add_chart(c1, 'L2')
    ws.add_chart(c2, 'L20')
    return ws

def graph_im(ws, _len):
    
    #-----------------------------------------------|
    # create chart
    c1 = LineChart()
    c1.title = 'Incoming Onnet'
    c1.y_axis.title = 'Minutes'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=2, min_row=4, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=4, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False

    # add chart to worksheet
    ws.add_chart(c1, 'I2')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'Incoming Offnet'
    c1.y_axis.title = 'Minutes'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=4, min_row=4, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=4, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False

    # add chart to worksheet
    ws.add_chart(c1, 'I18')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'Incoming PSTN'
    c1.y_axis.title = 'Minutes'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=6, min_row=4, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=4, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False

    # add chart to worksheet
    ws.add_chart(c1, 'I34')
    
    return ws

def graph_sd(ws, _len):
    
    #-----------------------------------------------|
    # create chart
    c1 = LineChart()
    c1.title = 'Global'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=3, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'A38')
    
    #-----------------------------------------------| 
    
    # create chart
    c1 = LineChart()
    c1.title = 'Moa Day ($7)'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=6, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'I38')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'Moa 2 Days ($14)'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=9, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'Q38')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'Moa 3 Days ($21)'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=12, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'A56')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'Moa Week ($42)'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=15, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'I56')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'Moa Month ($500)'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=18, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'Q56')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'D6 ($6)'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=21, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'A74')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'Hour Data ($10)'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=24, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'I74')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'D15 ($15)'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=27, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'Q74')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'D20 ($20)'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=33, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'A92')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'Movie Night ($35)'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=36, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'I92')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'Week Data ($50)'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=39, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'Q92')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'D90 ($90)'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=42, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'A110')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'D220 ($220)'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=45, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'I110')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'D500 ($500)'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=48, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'Q110')
    
    #-----------------------------------------------|
    
    # create chart
    c1 = LineChart()
    c1.title = 'Roaming Bundle'
    c1.y_axis.title = 'Amount (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=51, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'A128')
    
    #-----------------------------------------------|
    
    return ws

def graph_rbs(ws,_len):
    
    # create chart
    c1 = LineChart()
    c1.title = 'Total Recharge (SBD)'
    c1.y_axis.title = 'Recharge (SBD)'
    c1.x_axis.title = 'Date'
    
    # customize chart
    c1.style=2
    c1.y_axis.crossAx = 500
    c1.x_axis = DateAxis(crossAx=100)
    c1.x_axis.number_format = 'yyyy-mm-dd'
    c1.x_axis.majorTimeUnit = "days"
    c1.legend = None
    
    # y axis data
    data = Reference(ws, min_col=3, min_row=3, max_row=_len)
    c1.add_data(data)
    
    # x axis data
    dates = Reference(ws, min_col=1, min_row=3, max_row=_len)
    c1.set_categories(dates)
    
    # set chart width
    c1.width = 20
    
    # Style chart
    s1 = c1.series[0]
    s1.marker.symbol = 'circle'
    s1.marker.graphicalProperties.solidFill = 'FF0000' 
    s1.marker.graphicalProperties.line.solidFill = 'DDEBF7' 
    s1.graphicalProperties.line.noFill = False
    
    # add chart to worksheet
    ws.add_chart(c1, 'F2')
    
    return ws